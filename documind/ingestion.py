import io
import os
import tempfile
from pathlib import Path

import streamlit as st
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader


def _text_doc(text, filename, page=1):
    text = (text or "").strip()
    if not text:
        return []
    return [Document(page_content=text, metadata={"source": filename, "page": page})]


def load_documents(uploaded):
    name = uploaded.name
    suffix = Path(name).suffix.lower()
    data = uploaded.getvalue()
    docs = []

    if suffix == ".pdf":
        temp_path = None
        try:
            with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
                tmp.write(data)
                temp_path = tmp.name
            pages = PyPDFLoader(temp_path).load()
            for page in pages:
                page.metadata["source"] = name
                page.metadata["page"] = int(page.metadata.get("page", 0)) + 1
            docs = pages
        finally:
            if temp_path and os.path.exists(temp_path):
                os.remove(temp_path)
        # If PDF has no extractable text, attempt OCR if available.
        if not any(d.page_content.strip() for d in docs):
            try:
                import fitz
                from PIL import Image
                import pytesseract
                pdf = fitz.open(stream=data, filetype="pdf")
                for i, page in enumerate(pdf):
                    pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
                    image = Image.open(io.BytesIO(pix.tobytes("png")))
                    text = pytesseract.image_to_string(image)
                    docs.extend(_text_doc(text, name, i + 1))
            except Exception:
                pass

    elif suffix == ".docx":
        from docx import Document as DocxDocument
        docx = DocxDocument(io.BytesIO(data))
        text = "\n".join(p.text for p in docx.paragraphs)
        for i, table in enumerate(docx.tables, 1):
            text += "\n" + "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        docs = _text_doc(text, name)

    elif suffix in (".txt", ".md"):
        text = data.decode("utf-8", errors="replace")
        docs = _text_doc(text, name)

    elif suffix == ".epub":
        from ebooklib import epub
        from bs4 import BeautifulSoup
        with tempfile.NamedTemporaryFile(delete=False, suffix=".epub") as tmp:
            tmp.write(data)
            temp_path = tmp.name
        try:
            book = epub.read_epub(temp_path)
            section = 0
            for item in book.get_items():
                if item.get_type() == 9:  # ebooklib.ITEM_DOCUMENT
                    section += 1
                    soup = BeautifulSoup(item.get_content(), "html.parser")
                    docs.extend(_text_doc(soup.get_text(" ", strip=True), name, section))
        finally:
            os.remove(temp_path)

    elif suffix in (".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"):
        from PIL import Image
        import pytesseract
        image = Image.open(io.BytesIO(data))
        docs = _text_doc(pytesseract.image_to_string(image), name, 1)

    if not docs or not any(d.page_content.strip() for d in docs):
        raise ValueError(
            f"No readable text found in {name}. Scanned files require Tesseract OCR; "
            "check that it is installed and that the document is readable."
        )
    for doc in docs:
        doc.metadata.setdefault("source", name)
        doc.metadata.setdefault("page", 1)
    return docs


def chunk_documents(documents):
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=900, chunk_overlap=150,
        add_start_index=True, separators=["\n\n", "\n", ". ", " ", ""]
    )
    return splitter.split_documents(documents)

